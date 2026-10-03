"""results_table.py — Quản lý lưu trữ kết quả JSON và điền bảng Excel.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
import os
from pathlib import Path
import openpyxl


class SafeEncoder(json.JSONEncoder):
    def default(self, obj):
        if hasattr(obj, "item"):
            return obj.item()
        if hasattr(obj, "tolist"):
            return obj.tolist()
        return super().default(obj)


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    os.makedirs(results_dir, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    out_path = os.path.join(results_dir, f"{exp_id}.json")

    clean_dict = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(clean_dict, f, indent=2, ensure_ascii=False, cls=SafeEncoder)

    return out_path


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p = Path(results_dir)
    if not p.exists():
        return []
    res_files = sorted(p.glob("*.json"))
    results = []
    for f in res_files:
        with open(f, "r", encoding="utf-8") as fp:
            results.append(json.load(fp))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result["cfg"]
    summary = result.get("summary", {})
    exp_id = cfg["exp_id"]

    hidden_val = cfg.get("hidden", (256, 128))
    if isinstance(hidden_val, (list, tuple)):
        hidden_str = "-".join(str(h) for h in hidden_val)
    else:
        hidden_str = str(hidden_val)

    opt_raw = str(cfg.get("optimizer", "")).lower().strip()
    opt_map = {
        "sgd": "SGD",
        "sgd_momentum": "SGD+momentum",
        "adam": "Adam",
        "adamw": "AdamW",
    }
    opt_str = opt_map.get(opt_raw, cfg.get("optimizer", ""))

    clip_val = cfg.get("clip_norm")
    clip_str = "none" if (clip_val is None or str(clip_val).lower() == "none") else str(clip_val)

    row = dict(
        exp_id=exp_id,
        group=cfg.get("group", ""),
        description=cfg.get("description", ""),
        loss=cfg.get("loss", "ce").upper(),
        optimizer=opt_str,
        lr=cfg.get("lr", 0.0),
        weight_decay=cfg.get("weight_decay", 0.0),
        batch=cfg.get("batch", 512),
        epochs=cfg.get("epochs", 20),
        hidden=hidden_str,
        dropout=cfg.get("dropout", 0.0),
        clip_norm=clip_str,
        precision=cfg.get("precision", "fp32"),
        init=cfg.get("init", "he"),
        seed=cfg.get("seed", 1),
        step0_loss=round(float(summary.get("step0_loss", 0.0)), 4),
        best_val_loss=round(float(summary.get("best_val_loss", 0.0)), 4),
        best_epoch=int(summary.get("best_epoch", 1)),
        final_train_loss=round(float(summary.get("final_train_loss", 0.0)), 4),
        final_val_loss=round(float(summary.get("final_val_loss", 0.0)), 4),
        val_acc=round(float(summary.get("val_acc", 0.0)), 4),
        val_macro_f1=round(float(summary.get("val_macro_f1", 0.0)), 4),
        time_per_epoch_s=round(float(summary.get("time_per_epoch_s", 0.0)), 2),
        peak_mem_MB=round(float(summary.get("peak_mem_MB", 0.0)), 1),
        diverged=bool(summary.get("diverged", False)),
        eval_acc=round(float(eval_scores["accuracy"]), 4) if (eval_scores and "accuracy" in eval_scores) else "",
        eval_macro_f1=round(float(eval_scores["macro_f1"]), 4) if (eval_scores and "macro_f1" in eval_scores) else "",
        figure_file=f"figures/{exp_id}.png",
        notes=notes or cfg.get("notes", ""),
    )
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    header_cols = {}
    for col_idx in range(1, ws.max_column + 1):
        cell_val = ws.cell(row=1, column=col_idx).value
        if cell_val is not None:
            header_cols[str(cell_val).strip()] = col_idx

    formula_columns = {
        "step0_gap_vs_lnC",
        "gap_val_minus_train",
        "delta_val_f1_vs_base",
        "beyond_noise",
    }

    start_row = 2
    for r_idx, row_data in enumerate(rows, start=start_row):
        for key, val in row_data.items():
            if key in header_cols and key not in formula_columns:
                ws.cell(row=r_idx, column=header_cols[key], value=val)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)

